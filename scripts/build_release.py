#!/usr/bin/env python3
"""Build the multimodal READI release from owner-approved XLSX sources."""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

EXPECTED_COUNTS = {"ko": 57, "en": 45}
XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def normalize_intensity(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    match = re.fullmatch(r"\s*([1-3])(?:단계)?\s*", str(value))
    if not match:
        raise ValueError(f"invalid intensity value: {value!r}")
    return int(match.group(1))


def parse_choices(value: Any) -> dict[str, str]:
    text = "" if value is None else str(value).strip()
    pattern = re.compile(r"(?ms)^\s*([1-4])[\).]\s*(.*?)(?=^\s*[1-4][\).]\s*|\Z)")
    choices = {match.group(1): " ".join(match.group(2).split()) for match in pattern.finditer(text)}
    if set(choices) != {"1", "2", "3", "4"} or any(not v for v in choices.values()):
        raise ValueError("choice block must contain four non-empty numbered choices")
    return choices


def split_english_prompt(value: Any) -> tuple[str, str]:
    lines = [line.strip() for line in str(value or "").splitlines() if line.strip()]
    if len(lines) < 2:
        raise ValueError("English question/utterance cell cannot be separated")
    return lines[0], " ".join(lines[1:]).strip()


def load_corrections(path: Path | None) -> dict[str, Any]:
    if path is None or not path.exists():
        return {"intensity": {"ko": {}, "en": {}}}
    corrections = json.loads(path.read_text(encoding="utf-8"))
    if not corrections.get("approved_by") or not corrections.get("approved_at"):
        raise ValueError("correction file requires approved_by and approved_at")
    return corrections


def image_map(workbook: Path) -> dict[int, tuple[str, bytes]]:
    """Return one-based worksheet row -> (extension, image bytes)."""
    with zipfile.ZipFile(workbook) as archive:
        drawing = ET.fromstring(archive.read("xl/drawings/drawing1.xml"))
        rels_root = ET.fromstring(archive.read("xl/drawings/_rels/drawing1.xml.rels"))
        rels = {
            rel.attrib["Id"]: rel.attrib["Target"]
            for rel in rels_root.findall(f"{{{PKG_REL}}}Relationship")
        }
        result = {}
        for anchor in list(drawing):
            start = anchor.find(f"{{{XDR}}}from")
            blip = anchor.find(f".//{{{A}}}blip")
            if start is None or blip is None:
                continue
            row = int(start.find(f"{{{XDR}}}row").text) + 1
            rel_id = blip.attrib[f"{{{R}}}embed"]
            target = posixpath.normpath(posixpath.join("xl/drawings", rels[rel_id]))
            suffix = Path(target).suffix.lower()
            if suffix not in {".png", ".jpg", ".jpeg"}:
                raise ValueError(f"unsupported image format: {suffix}")
            if row in result:
                raise ValueError(f"multiple images anchored to worksheet row {row}")
            result[row] = (".jpg" if suffix == ".jpeg" else suffix, archive.read(target))
        return result


def source_records(path: Path, language: str, corrections: dict[str, Any]):
    ws = load_workbook(path, read_only=False, data_only=True).active
    images = image_map(path)
    records = []
    if language == "ko":
        columns = {"source_id": 1, "question": 4, "utterance": 5, "choices": 6, "answer": 7, "intensity": 8}
    else:
        columns = {"source_id": 1, "reference": 2, "combined": 4, "choices": 5, "answer": 6, "intensity": 7}

    item_rows = [row for row in range(2, ws.max_row + 1) if ws.cell(row, 1).value is not None]
    if set(item_rows) != set(images):
        raise ValueError(f"{language}: image anchors and item rows differ")

    for public_index, row in enumerate(item_rows, 1):
        source_id = str(ws.cell(row, columns["source_id"]).value).strip()
        intensity = normalize_intensity(ws.cell(row, columns["intensity"]).value)
        if intensity is None:
            supplied = corrections.get("intensity", {}).get(language, {}).get(source_id)
            intensity = normalize_intensity(supplied)
        if intensity is None:
            raise ValueError(f"{language} source item {source_id}: missing intensity")

        if language == "ko":
            question = str(ws.cell(row, columns["question"]).value or "").strip()
            utterance = str(ws.cell(row, columns["utterance"]).value or "").strip()
            source_reference = None
        else:
            question, utterance = split_english_prompt(ws.cell(row, columns["combined"]).value)
            source_reference = str(ws.cell(row, columns["reference"]).value or "").strip() or None
        if not question or not utterance:
            raise ValueError(f"{language} source item {source_id}: empty question or utterance")

        answer_text = str(ws.cell(row, columns["answer"]).value or "").strip()
        match = re.fullmatch(r"([1-4])(?:\.0)?", answer_text)
        if not match:
            raise ValueError(f"{language} source item {source_id}: invalid answer")
        extension, image_bytes = images[row]
        item_id = f"readi_{language}_{public_index:03d}"
        record = {
            "item_id": item_id,
            "language": language,
            "image": f"images/{item_id}{extension}",
            "question": question,
            "utterance": utterance,
            "choices": parse_choices(ws.cell(row, columns["choices"]).value),
            "answer": int(match.group(1)),
            "isa_intensity": intensity,
        }
        if source_reference:
            record["source_reference"] = source_reference
        records.append((record, extension, image_bytes))
    return records


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--krisa", type=Path, required=True)
    parser.add_argument("--engisa", type=Path, required=True)
    parser.add_argument("--corrections", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--allow-count-mismatch", action="store_true", help="Audit only; never use for a public release")
    args = parser.parse_args()

    corrections = load_corrections(args.corrections)
    grouped = {
        "ko": source_records(args.krisa, "ko", corrections),
        "en": source_records(args.engisa, "en", corrections),
    }
    for language, records in grouped.items():
        if len(records) != EXPECTED_COUNTS[language] and not args.allow_count_mismatch:
            raise ValueError(
                f"{language}: workbook has {len(records)} items, paper reports {EXPECTED_COUNTS[language]}"
            )

    args.output_dir.mkdir(parents=True, exist_ok=True)
    image_dir = args.output_dir / "images"
    image_dir.mkdir(parents=True, exist_ok=True)
    all_records = []
    for language in ("ko", "en"):
        for record, extension, image_bytes in grouped[language]:
            output_image = args.output_dir / record["image"]
            output_image.write_bytes(image_bytes)
            all_records.append(record)
    data_path = args.output_dir / "readi.jsonl"
    data_path.write_text(
        "".join(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n" for record in all_records),
        encoding="utf-8",
    )
    manifest = {
        "format_version": "1.0.0",
        "release_ready": all(len(grouped[k]) == EXPECTED_COUNTS[k] for k in grouped),
        "counts": {k: len(v) for k, v in grouped.items()},
        "records": len(all_records),
        "files": {"readi.jsonl": sha256(data_path)},
    }
    for image in sorted(image_dir.iterdir()):
        manifest["files"][str(image.relative_to(args.output_dir))] = sha256(image)
    (args.output_dir / "manifest.json").write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(manifest | {"files": f"{len(manifest['files'])} hashed files"}, indent=2))


if __name__ == "__main__":
    main()

